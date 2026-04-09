from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app


@pytest.fixture()
def client(tmp_path):
    settings.storage_dir = str(tmp_path / "storage")
    settings.input_dir = str(tmp_path / "storage" / "inputs")
    settings.output_dir = str(tmp_path / "storage" / "outputs")
    settings.report_dir = str(tmp_path / "storage" / "reports")

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as c:
        yield c
    app.dependency_overrides.clear()


def _form_data():
    return {
        "faculty": "ФКТ",
        "department": "Кафедра ИС",
        "lab_title": "Тестирование API",
        "lab_number": "1",
        "student_name": "Иванов И.И.",
        "reviewer_name": "Петров П.П.",
        "discipline": "Программирование",
    }


def _valid_docx_bytes() -> bytes:
    doc = Document()
    long_text = ("Тестовый текст " * 260).strip()
    doc.add_paragraph(long_text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _invalid_by_rules_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Короткий текст")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_create_task_success_and_artifacts_available(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "completed"
    task_id = data["task_id"]

    status_resp = client.get(f"/tasks/{task_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["has_output"] is True
    assert status_data["has_report"] is True

    download_resp = client.get(f"/tasks/{task_id}/download")
    assert download_resp.status_code == 200
    assert (
        download_resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    report_resp = client.get(f"/tasks/{task_id}/report")
    assert report_resp.status_code == 200
    assert report_resp.headers["content-type"].startswith("application/json")


def test_create_task_rejects_non_docx(client: TestClient):
    files = {
        "file": (
            "bad.txt",
            b"plain text",
            "text/plain",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    assert "DOCX" in resp.json()["detail"]


def test_create_task_rejects_empty_docx(client: TestClient):
    files = {
        "file": (
            "empty.docx",
            b"",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    assert "пустой" in resp.json()["detail"].lower()


def test_create_task_rejects_corrupted_docx(client: TestClient):
    files = {
        "file": (
            "corrupted.docx",
            b"this is not a real docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 400
    assert "docx" in resp.json()["detail"].lower()


def test_create_task_fails_business_validation(client: TestClient):
    files = {
        "file": (
            "too_short.docx",
            _invalid_by_rules_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    resp = client.post("/tasks", data=_form_data(), files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "failed"
    assert data["report"]["status"] == "failed"
    assert data["report"]["errors"]


def test_get_unknown_task_returns_404(client: TestClient):
    resp = client.get("/tasks/not-existing-id")
    assert resp.status_code == 404


def test_list_tasks_history_for_frontend(client: TestClient):
    files = {
        "file": (
            "valid.docx",
            _valid_docx_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    create_resp = client.post("/tasks", data=_form_data(), files=files)
    assert create_resp.status_code == 200
    created_task_id = create_resp.json()["task_id"]

    history_resp = client.get("/tasks?limit=10")
    assert history_resp.status_code == 200
    body = history_resp.json()
    assert "items" in body
    assert isinstance(body["items"], list)
    assert len(body["items"]) >= 1

    first = body["items"][0]
    assert first["task_id"] == created_task_id
    assert "status" in first
    assert "original_filename" in first
    assert "created_at" in first
    assert "has_output" in first
    assert "has_report" in first
