import json
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.config import settings
from app.core.db import Base, get_db
from app.main import app


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("test")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _valid_processing_docx_bytes() -> bytes:
    doc = Document()
    long_text = ("Тестовый текст " * 260).strip()
    doc.add_paragraph(long_text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _client(tmp_path):
    settings.storage_dir = str(tmp_path / "storage")
    settings.input_dir = str(tmp_path / "storage" / "inputs")
    settings.output_dir = str(tmp_path / "storage" / "outputs")
    settings.report_dir = str(tmp_path / "storage" / "reports")
    settings.projects_dir = str(tmp_path / "storage" / "projects")

    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = TestingSessionLocal()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    return TestClient(app)


def test_upload_project_file_creates_project_structure(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert resp.status_code == 200
        body = resp.json()
        assert body["status"] == "uploaded"
        project_id = body["project_id"]

        project_dir = Path(settings.projects_dir) / project_id
        assert (project_dir / "input").exists()
        assert (project_dir / "images").exists()
        assert (project_dir / "output").exists()
        metadata_path = project_dir / "metadata.json"
        assert metadata_path.exists()

        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        assert metadata["db_snapshot"]["project_id"] == project_id
        assert metadata["db_snapshot"]["status"] == "uploaded"
        assert metadata["metadata"]["source_filename"] == "source.docx"
        assert metadata["metadata"]["files"][0]["name"] == "source.docx"


def test_upload_project_file_rejects_unsupported_extension(tmp_path):
    with _client(tmp_path) as client:
        files = {"file": ("notes.txt", b"text", "text/plain")}
        resp = client.post("/projects/upload", files=files)
        assert resp.status_code == 400
        assert "Поддерживаются только форматы" in resp.json()["message"]


def test_download_project_result_returns_latest_completed_output(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        task_data = {
            "user_id": "user-1",
            "project_id": project_id,
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        task_files = {
            "file": (
                "valid.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        task_resp = client.post("/tasks", data=task_data, files=task_files)
        assert task_resp.status_code == 200

        download_resp = client.get(f"/projects/{project_id}/download?user_id=user-1")
        assert download_resp.status_code == 200
        assert (
            download_resp.headers["content-type"]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


def test_download_project_result_forbidden_for_other_user(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        forbidden_resp = client.get(f"/projects/{project_id}/download?user_id=user-2")
        assert forbidden_resp.status_code == 403


def test_delete_project_removes_project_tasks_and_project_folder(tmp_path):
    with _client(tmp_path) as client:
        upload_files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=upload_files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        task_data = {
            "user_id": "user-1",
            "project_id": project_id,
            "faculty": "ФКТ",
            "department": "Кафедра ИС",
            "student_group": "БПИ-01",
            "lab_title": "Тестирование API",
            "lab_number": "1",
            "student_name": "Иванов И.И.",
            "reviewer_name": "Петров П.П.",
            "discipline": "Программирование",
        }
        task_files = {
            "file": (
                "valid.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        task_resp = client.post("/tasks", data=task_data, files=task_files)
        assert task_resp.status_code == 200
        task_id = task_resp.json()["task_id"]

        delete_resp = client.delete(f"/projects/{project_id}?user_id=user-1")
        assert delete_resp.status_code == 200
        assert delete_resp.json()["status"] == "deleted"

        status_resp = client.get(f"/projects/{project_id}")
        assert status_resp.status_code == 404

        task_status_resp = client.get(f"/tasks/{task_id}")
        assert task_status_resp.status_code == 404

        project_dir = Path(settings.projects_dir) / project_id
        assert not project_dir.exists()


def test_delete_project_forbidden_for_other_user(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.docx",
                _valid_processing_docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        forbidden_resp = client.delete(f"/projects/{project_id}?user_id=user-2")
        assert forbidden_resp.status_code == 403


def test_project_analyze_and_get_analysis(tmp_path):
    with _client(tmp_path) as client:
        files = {
            "file": (
                "source.pdf",
                b"%PDF-1.4 test content",
                "application/pdf",
            )
        }
        upload_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=files)
        assert upload_resp.status_code == 200
        project_id = upload_resp.json()["project_id"]

        analyze_resp = client.post(f"/projects/{project_id}/analyze?user_id=user-1")
        assert analyze_resp.status_code == 200
        analyze_body = analyze_resp.json()
        assert analyze_body["status"] == "ready"
        assert "analysis" in analyze_body
        assert analyze_body["analysis"]["source_extension"] == ".pdf"

        get_resp = client.get(f"/projects/{project_id}/analysis?user_id=user-1")
        assert get_resp.status_code == 200
        get_body = get_resp.json()
        assert get_body["analysis"]["model"] == "mock-ml-v1"

        metadata_path = Path(settings.projects_dir) / project_id / "metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        assert metadata["metadata"]["analysis"]["model"] == "mock-ml-v1"
        assert len(metadata["metadata"]["ml_suggestions"]) > 0


def test_upload_additional_files_to_existing_project_separates_input_and_images(tmp_path):
    with _client(tmp_path) as client:
        create_files = {
            "file": (
                "base.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        create_resp = client.post("/projects/upload", data={"user_id": "user-1"}, files=create_files)
        assert create_resp.status_code == 200
        project_id = create_resp.json()["project_id"]

        docx_files = {
            "file": (
                "extra.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        }
        add_docx_resp = client.post(
            f"/projects/{project_id}/files",
            data={"user_id": "user-1"},
            files=docx_files,
        )
        assert add_docx_resp.status_code == 200
        assert add_docx_resp.json()["file_type"] == "input"

        image_files = {"file": ("scan.jpg", b"jpeg-bytes", "image/jpeg")}
        add_image_resp = client.post(
            f"/projects/{project_id}/files",
            data={"user_id": "user-1"},
            files=image_files,
        )
        assert add_image_resp.status_code == 200
        assert add_image_resp.json()["file_type"] == "image"

        project_dir = Path(settings.projects_dir) / project_id
        assert (project_dir / "input" / "extra.docx").exists()
        assert (project_dir / "images" / "scan.jpg").exists()

        metadata_path = project_dir / "metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        file_names = {item["name"] for item in metadata["metadata"]["files"]}
        assert "extra.docx" in file_names
        assert "scan.jpg" in file_names
