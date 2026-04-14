from __future__ import annotations

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

from app.schemas.blocks import ParagraphBlock, TableBlock, ImageBlock
from app.core.config import settings


def apply_gost_page_settings(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)


def set_run_font(run) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)


def format_paragraph_common(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def add_text_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    run = p.add_run(text)
    set_run_font(run)
    format_paragraph_common(p)


def add_centered_paragraph(doc: Document, text: str, bold: bool = False, line_spacing: float = 1.0, space_before: float = 0, space_after: float = 0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    set_run_font(run)


def add_right_paragraph(doc: Document, text: str, line_spacing: float = 1.0, space_before: float = 0) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run)


def add_empty_line(doc: Document, line_spacing: float = 1.0) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)


def build_title_page(doc: Document, payload: dict) -> None:
    add_centered_paragraph(doc, settings.university_name, bold=True, line_spacing=1.0)
    add_centered_paragraph(doc, f"Факультет «{payload['faculty']}»", line_spacing=1.0)
    add_centered_paragraph(doc, f"Кафедра «{payload['department']}»", line_spacing=1.0)

    add_empty_line(doc)
    add_empty_line(doc)

    add_centered_paragraph(doc, f"Отчет по лабораторной работе № {payload['lab_number']}", bold=True, line_spacing=1.0)
    add_centered_paragraph(doc, f"по дисциплине «{payload['discipline']}»", line_spacing=1.0)
    add_centered_paragraph(doc, "на тему:", line_spacing=1.0)
    add_centered_paragraph(doc, f"«{payload['lab_title']}»", line_spacing=1.0)

    add_empty_line(doc)
    add_empty_line(doc)
    add_empty_line(doc)

    add_right_paragraph(doc, f"Выполнил: студент группы {payload['student_group']}", line_spacing=1.0)
    add_right_paragraph(doc, f"{payload['student_name']}", line_spacing=1.0)
    add_right_paragraph(doc, f"Проверил:", line_spacing=1.0)
    add_right_paragraph(doc, f"{payload['reviewer_name']}", line_spacing=1.0)

    add_empty_line(doc)
    add_empty_line(doc)
    add_empty_line(doc)

    add_centered_paragraph(doc, f"{settings.city}, {settings.year}", space_before=180)
    doc.add_section(WD_SECTION.NEW_PAGE)


def add_table_block(doc: Document, block: TableBlock) -> None:
    if block.caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(block.caption)
        set_run_font(run)
        p.paragraph_format.line_spacing = 1.5

    if not block.rows:
        return

    rows_count = len(block.rows)
    cols_count = max(len(row) for row in block.rows) if block.rows else 1

    table = doc.add_table(rows=rows_count, cols=cols_count)
    table.style = "Table Grid"

    for i, row in enumerate(block.rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)


def add_image_block(doc: Document, block: ImageBlock) -> None:
    if block.image_path:
        try:
            doc.add_picture(block.image_path, width=Cm(12))
        except Exception:
            add_text_paragraph(doc, "[Не удалось вставить изображение]")
    else:
        add_text_paragraph(doc, "[Изображение отсутствует]")

    if block.caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(block.caption)
        set_run_font(run)
        p.paragraph_format.line_spacing = 1.5


def build_output_document(blocks: list, payload: dict, output_path: str) -> None:
    doc = Document()
    apply_gost_page_settings(doc)

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(14)

    build_title_page(doc, payload)

    for block in blocks:
        if isinstance(block, ParagraphBlock):
            if block.text.strip():
                add_text_paragraph(doc, block.text)

        elif isinstance(block, TableBlock):
            add_table_block(doc, block)

        elif isinstance(block, ImageBlock):
            add_image_block(doc, block)

    doc.save(output_path)