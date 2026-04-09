from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm
from tempfile import TemporaryDirectory

from app.services.extractor import extract_blocks
from app.services.title_detector import remove_existing_title_page
from app.services.captions import attach_captions, ensure_generated_captions
from app.services.validator import validate_blocks
from app.services.renderer import build_output_document
from app.core.config import settings


FIGURE_RE = re.compile(r'^\s*Рисунок\s+\d+[\s\S]*$', re.IGNORECASE)
TABLE_RE = re.compile(r'^\s*Таблица\s+\d+[\s\S]*$', re.IGNORECASE)


def _set_default_font(run) -> None:
    run.font.name = settings.font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), settings.font_name)
    run.font.size = Pt(settings.font_size_pt)


def _style_paragraph(paragraph, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=True, bold=False, font_size=None):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = settings.line_spacing
    fmt.first_line_indent = Cm(settings.first_line_indent_cm) if first_line_indent else Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    paragraph.alignment = align

    for run in paragraph.runs:
        _set_default_font(run)
        run.bold = bold
        if font_size:
            run.font.size = Pt(font_size)


def _set_page_layout(doc: DocumentType) -> None:
    for section in doc.sections:
        section.top_margin = Cm(settings.margin_top_cm)
        section.bottom_margin = Cm(settings.margin_bottom_cm)
        section.left_margin = Cm(settings.margin_left_cm)
        section.right_margin = Cm(settings.margin_right_cm)


def _clear_document_body(doc: DocumentType) -> None:
    body = doc._element.body
    for child in list(body):
        body.remove(child)


def _add_page_break(paragraph) -> None:
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def _insert_paragraph_before(paragraph, text: str = ''):
    new_p = OxmlElement('w:p')
    paragraph._p.addprevious(new_p)
    from docx.text.paragraph import Paragraph
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _insert_paragraph_after(paragraph, text: str = ''):
    new_p = OxmlElement('w:p')
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def _is_image_paragraph(paragraph) -> bool:
    xml = paragraph._p.xml
    return 'graphicData' in xml or 'pic:pic' in xml or 'a:blip' in xml


def _has_table_caption_before(doc: DocumentType, table_index: int) -> bool:
    body_items = list(doc.element.body.iterchildren())
    current_tbl_seen = -1
    for i, item in enumerate(body_items):
        if item.tag.endswith('}tbl'):
            current_tbl_seen += 1
            if current_tbl_seen == table_index:
                if i > 0 and body_items[i - 1].tag.endswith('}p'):
                    text = ''.join(body_items[i - 1].itertext()).strip()
                    return bool(TABLE_RE.match(text))
                return False
    return False


def _iter_body_paragraphs(doc: DocumentType):
    for p in doc.paragraphs:
        yield p


def build_title_page(doc: DocumentType, payload: dict) -> None:
    title_blocks = [
        (settings.university_name, WD_ALIGN_PARAGRAPH.CENTER, True, 14, False),
        (payload['faculty'], WD_ALIGN_PARAGRAPH.CENTER, False, 14, False),
        (payload['department'], WD_ALIGN_PARAGRAPH.CENTER, False, 14, False),
        ('', WD_ALIGN_PARAGRAPH.CENTER, False, 14, False),
        ('ЛАБОРАТОРНАЯ РАБОТА', WD_ALIGN_PARAGRAPH.CENTER, True, 16, False),
        (f'№ {payload["lab_number"]}', WD_ALIGN_PARAGRAPH.CENTER, True, 14, False),
        (payload['lab_title'], WD_ALIGN_PARAGRAPH.CENTER, True, 14, False),
        ('', WD_ALIGN_PARAGRAPH.CENTER, False, 14, False),
    ]

    for text, align, bold, size, first_indent in title_blocks:
        p = doc.add_paragraph()
        run = p.add_run(text)
        _set_default_font(run)
        run.bold = bold
        run.font.size = Pt(size)
        p.alignment = align
        p.paragraph_format.line_spacing = settings.line_spacing
        p.paragraph_format.first_line_indent = Cm(0 if not first_indent else settings.first_line_indent_cm)

    info_lines = [
        f'Выполнил: {payload["student_name"]}',
        f'Проверил: {payload["reviewer_name"]}',
    ]
    for line in info_lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        _set_default_font(run)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = settings.line_spacing
        p.paragraph_format.first_line_indent = Cm(0)

    for _ in range(8):
        doc.add_paragraph('')

    footer_p = doc.add_paragraph()
    footer_run = footer_p.add_run(f'{settings.city}, {settings.year}')
    _set_default_font(footer_run)
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_page_break(doc.add_paragraph())


def append_source_content(target_doc: DocumentType, source_doc: DocumentType) -> None:
    for paragraph in source_doc.paragraphs:
        new_p = target_doc.add_paragraph()
        for run in paragraph.runs:
            new_run = new_p.add_run(run.text)
            _set_default_font(new_run)
            new_run.bold = bool(run.bold)
            new_run.italic = bool(run.italic)
            new_run.underline = bool(run.underline)
        if not paragraph.runs and paragraph.text:
            new_run = new_p.add_run(paragraph.text)
            _set_default_font(new_run)

        _style_paragraph(new_p)

        if _is_image_paragraph(paragraph):
            # Прямое копирование embedded-изображений через python-docx ограничено.
            # Для MVP оставляем маркер, чтобы не терять место изображения при создании нового файла.
            if not new_p.text.strip():
                marker = new_p.add_run('[Изображение из исходного документа]')
                _set_default_font(marker)
                marker.italic = True
                new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for table in source_doc.tables:
        rows = len(table.rows)
        cols = len(table.columns)
        new_table = target_doc.add_table(rows=rows, cols=cols)
        new_table.style = 'Table Grid'
        for i in range(rows):
            for j in range(cols):
                cell_text = table.cell(i, j).text
                new_table.cell(i, j).text = cell_text
                for p in new_table.cell(i, j).paragraphs:
                    _style_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=False)


def normalize_body_text(doc: DocumentType) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if FIGURE_RE.match(text) or TABLE_RE.match(text):
            _style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=False)
            continue
        if text == 'ЛАБОРАТОРНАЯ РАБОТА' or text.startswith('№ '):
            continue
        _style_paragraph(paragraph)


def add_missing_table_captions(doc: DocumentType, report: dict) -> None:
    body_items = list(doc.element.body.iterchildren())
    table_idx = 0
    table_number = 1
    for i, item in enumerate(body_items):
        if item.tag.endswith('}tbl'):
            has_caption = False
            if i > 0 and body_items[i - 1].tag.endswith('}p'):
                text = ''.join(body_items[i - 1].itertext()).strip()
                has_caption = bool(TABLE_RE.match(text))
            if not has_caption:
                from docx.text.paragraph import Paragraph
                prev_or_self = item
                pxml = OxmlElement('w:p')
                prev_or_self.addprevious(pxml)
                p = Paragraph(pxml, doc)
                run = p.add_run(f'Таблица {table_number}')
                _set_default_font(run)
                _style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=False)
                report['fixes'].append(f'Добавлена подпись для таблицы {table_number}.')
            table_idx += 1
            table_number += 1


def add_missing_figure_captions(doc: DocumentType, report: dict) -> None:
    figure_number = 1
    paragraphs = list(doc.paragraphs)
    for idx, paragraph in enumerate(paragraphs):
        if not _is_image_paragraph(paragraph) and '[Изображение из исходного документа]' not in paragraph.text:
            continue

        next_text = paragraphs[idx + 1].text.strip() if idx + 1 < len(paragraphs) else ''
        if FIGURE_RE.match(next_text):
            figure_number += 1
            continue

        new_p = _insert_paragraph_after(paragraph, f'Рисунок {figure_number}')
        _style_paragraph(new_p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=False)
        report['fixes'].append(f'Добавлена подпись для рисунка {figure_number}.')
        figure_number += 1


def process_document(input_path: str, output_path: str, title_payload: dict) -> dict:
    with TemporaryDirectory() as temp_dir:
        blocks = extract_blocks(input_path, temp_dir)
        blocks = remove_existing_title_page(blocks)
        blocks = attach_captions(blocks)
        blocks = ensure_generated_captions(blocks)

        validation = validate_blocks(blocks)

        if not validation["is_valid"]:
            return {
                "success": False,
                "warnings": validation["warnings"],
                "errors": validation["errors"],
                "changes": [],
                "validation": validation,
            }

        build_output_document(blocks, title_payload, output_path)

        return {
            "success": True,
            "warnings": validation["warnings"],
            "errors": validation["errors"],
            "changes": [
                "Создан новый титульный лист по шаблону.",
                "Старый титульный лист исключён из итогового документа.",
                "Документ пересобран в новый DOCX.",
                "Выполнена генерация и нормализация подписей рисунков и таблиц.",
                "Применено фиксированное оформление по ГОСТ-профилю.",
            ],
            "validation": validation,
        }
