from __future__ import annotations

from copy import deepcopy
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


def _paragraph_index(target_id: str) -> int:
    return int(target_id.split('_')[1])


def _run_indexes(target_id: str) -> tuple[int, int]:
    _, p_idx, r_idx = target_id.split('_')
    return int(p_idx), int(r_idx)


def _section_index(target_id: str) -> int:
    return int(target_id.split('_')[1])


def inverse_fix(fix: dict) -> dict:
    inverse = deepcopy(fix)
    inverse['fix_id'] = f"{fix['fix_id']}_undo"
    inverse['old_value'], inverse['new_value'] = fix.get('new_value'), fix.get('old_value')
    return inverse


def apply_fix_to_doc(doc, fix: dict):
    action = fix['action']

    if action == 'set_first_line_indent':
        idx = _paragraph_index(fix['target_id'])
        doc.paragraphs[idx].paragraph_format.first_line_indent = Cm(float(fix['new_value']))

    elif action == 'set_line_spacing':
        idx = _paragraph_index(fix['target_id'])
        doc.paragraphs[idx].paragraph_format.line_spacing = float(fix['new_value'])

    elif action == 'set_space_before':
        idx = _paragraph_index(fix['target_id'])
        doc.paragraphs[idx].paragraph_format.space_before = Pt(float(fix['new_value']))

    elif action == 'set_space_after':
        idx = _paragraph_index(fix['target_id'])
        doc.paragraphs[idx].paragraph_format.space_after = Pt(float(fix['new_value']))

    elif action == 'set_alignment_justify':
        idx = _paragraph_index(fix['target_id'])
        doc.paragraphs[idx].alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    elif action == 'set_font_name':
        p_idx, r_idx = _run_indexes(fix['target_id'])
        doc.paragraphs[p_idx].runs[r_idx].font.name = fix['new_value']

    elif action == 'set_font_size':
        p_idx, r_idx = _run_indexes(fix['target_id'])
        doc.paragraphs[p_idx].runs[r_idx].font.size = Pt(float(fix['new_value']))

    elif action == 'remove_trailing_period':
        idx = _paragraph_index(fix['target_id'])
        paragraph = doc.paragraphs[idx]
        text = paragraph.text.rstrip('.')
        if paragraph.runs:
            paragraph.runs[0].text = text
            for run in paragraph.runs[1:]:
                run.text = ''

    elif action == 'set_left_margin':
        idx = _section_index(fix['target_id'])
        doc.sections[idx].left_margin = Cm(float(fix['new_value']))

    elif action == 'set_right_margin':
        idx = _section_index(fix['target_id'])
        doc.sections[idx].right_margin = Cm(float(fix['new_value']))

    elif action == 'set_top_margin':
        idx = _section_index(fix['target_id'])
        doc.sections[idx].top_margin = Cm(float(fix['new_value']))

    elif action == 'set_bottom_margin':
        idx = _section_index(fix['target_id'])
        doc.sections[idx].bottom_margin = Cm(float(fix['new_value']))

    else:
        raise ValueError(f'Unsupported fix action: {action}')

    return doc
